# -*- coding: utf-8 -*-
"""Formula generators and styling helpers for the design workbook."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
C_IN     = Font(name=FONT, size=10, color="0000FF", bold=True)   # user input
C_CALC   = Font(name=FONT, size=10)                              # formula
C_LINK   = Font(name=FONT, size=10, color="008000")              # cross-sheet
C_HEAD   = Font(name=FONT, size=11, bold=True, color="FFFFFF")
C_SEC    = Font(name=FONT, size=10, bold=True, color="1F3864")
C_TITLE  = Font(name=FONT, size=14, bold=True, color="1F3864")
C_NOTE   = Font(name=FONT, size=9, italic=True, color="595959")
C_RES    = Font(name=FONT, size=10, bold=True, color="C00000")

F_IN     = PatternFill("solid", fgColor="FFF2CC")
F_HEAD   = PatternFill("solid", fgColor="1F3864")
F_SEC    = PatternFill("solid", fgColor="D9E2F3")
F_RES    = PatternFill("solid", fgColor="FCE4D6")
F_OK     = PatternFill("solid", fgColor="E2EFDA")

THIN = Side(style="thin", color="BFBFBF")
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ------------------------------------------------------------------ constants
RU      = "8.31446261815324"
M_H2O   = "0.018015268"
M_CO2   = "0.0440095"
TC_W, PC_W = "647.096", "22064000"
RV      = "461.5228"          # R_u / M_H2O

def PSAT(T):
    """Wagner & Pruss (IAPWS-95) saturation pressure [Pa] for a temperature cell T [K]."""
    t = f"(1-{T}/{TC_W})"
    return (f"{PC_W}*EXP({TC_W}/{T}*(-7.85951783*{t}+1.84408259*{t}^1.5"
            f"-11.7866497*{t}^3+22.6807411*{t}^3.5-15.9618719*{t}^4"
            f"+1.80122502*{t}^7.5))")

def LV(T):
    """Enthalpy of vaporisation of water [J/kg] at T [K] (fit to IAPWS-95)."""
    return f"1000*(2500.93-2.3666*({T}-273.15)-0.0016*({T}-273.15)^2)"

def DPSATDT(psat_cell, T):
    """dp_sat/dT via Clausius-Clapeyron [Pa/K] -- used only as the Newton slope."""
    return f"{psat_cell}*{LV(T)}/({RV}*{T}^2)"

def TDEW_SEED(pv):
    """Alduchov & Eskridge (1996) inverse Magnus seed [K]."""
    return f"273.15+243.04*LN({pv}/610.94)/(17.625-LN({pv}/610.94))"

def BCO2(T):
    """2nd virial coefficient of CO2 [m3/mol]; fit to Span & Wagner (1996)."""
    t = f"({T}/100)"
    return f"0.000001*(75.188632-565.620467/{t}+601.464888/{t}^2-2018.286032/{t}^3)"

def BWW(T):
    """2nd virial coefficient of water vapour [m3/mol]; fit to IAPWS-95."""
    return (f"0.000001*(19202.081-20032490/{T}+7075968700/{T}^2"
            f"-870743060000/{T}^3)")

def VWL(T):
    """Molar volume of liquid water [m3/mol]; Kell (1975)."""
    t = f"({T}-273.15)"
    return (f"{M_H2O}/((999.83952+16.945176*{t}-0.0079870401*{t}^2"
            f"-0.000046170461*{t}^3+0.00000010556302*{t}^4"
            f"-0.00000000028054253*{t}^5)/(1+0.016879850*{t}))")

def ZCO2(p, T):
    return f"1+{BCO2(T)}*{p}/({RU}*{T})"

def CP0CO2(T):
    """Ideal-gas molar cp of CO2 [J/(mol K)]; fit to Span & Wagner (1996)."""
    return f"20.823075+0.059235957*{T}-0.00000051159127*{T}^2-0.000000048979736*{T}^3"

def CP0H2O(T):
    """Ideal-gas molar cp of water vapour [J/(mol K)]; Shomate / NIST-JANAF."""
    t = f"({T}/1000)"
    return f"30.0920+6.832514*{t}+6.793435*{t}^2-2.534480*{t}^3+0.082139/{t}^2"

def MUCO2(T):
    """Dilute-gas viscosity of CO2 [Pa s]; Fenghour, Wakeham & Vesovic (1998)."""
    l = f"LN({T}/251.196)"
    return (f"0.00000100697*SQRT({T})/EXP(0.235156-0.491266*{l}"
            f"+0.05211155*{l}^2+0.05347906*{l}^3-0.01537102*{l}^4)")

def KCO2(T):
    """Dilute-gas thermal conductivity of CO2 [W/(m K)]; fit to Huber et al. (2016)."""
    return f"0.001*(-2.8088692+0.053424539*{T}+0.000038923913*{T}^2)"

def CSTAR(g):
    """Perfect-gas critical flow function (ISO 9300)."""
    return f"SQRT({g}*(2/({g}+1))^(({g}+1)/({g}-1)))"

def RCRIT(g):
    """Critical (choking) pressure ratio p*/p0."""
    return f"(2/({g}+1))^({g}/({g}-1))"

# ------------------------------------------------------------------- layout
def title(ws, row, text, sub=""):
    ws.cell(row, 1, text).font = C_TITLE
    if sub:
        ws.cell(row+1, 1, sub).font = C_NOTE
    return row + (3 if sub else 2)

def section(ws, row, text, width=9):
    for c in range(1, width+1):
        ws.cell(row, c).fill = F_SEC
    ws.cell(row, 1, text).font = C_SEC
    return row + 1

def hdr(ws, row, labels, col0=1):
    for i, l in enumerate(labels):
        c = ws.cell(row, col0+i, l)
        c.font, c.fill, c.border = C_HEAD, F_HEAD, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return row + 1

def inp(ws, row, label, value, unit="", note="", fmt="0.000"):
    ws.cell(row, 1, label).font = C_CALC
    c = ws.cell(row, 2, value); c.font, c.fill, c.border = C_IN, F_IN, BOX
    c.number_format = fmt
    ws.cell(row, 3, unit).font = C_CALC
    if note: ws.cell(row, 4, note).font = C_NOTE
    return row + 1

def out(ws, row, label, formula, unit="", note="", fmt="0.0000", res=False):
    ws.cell(row, 1, label).font = C_CALC
    c = ws.cell(row, 2, formula)
    c.font = C_RES if res else C_CALC
    if res: c.fill = F_RES
    c.border, c.number_format = BOX, fmt
    ws.cell(row, 3, unit).font = C_CALC
    if note: ws.cell(row, 4, note).font = C_NOTE
    return row + 1

def note(ws, row, text, col=1):
    ws.cell(row, col, text).font = C_NOTE
    return row + 1

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def apply_font(ws, first_row=1):
    for r in ws.iter_rows(min_row=first_row):
        for c in r:
            if c.font is None or c.font.name is None:
                c.font = C_CALC

def DBDT_CO2(T):
    """Analytic dB/dT for the CO2 2nd virial fit [m3/(mol K)]."""
    t = f"({T}/100)"
    return f"0.00000001*(565.620467/{t}^2-1202.929776/{t}^3+6054.858096/{t}^4)"

def MUJT_CO2(T):
    """Joule-Thomson coefficient [K/Pa], zero-density limit (T dB/dT - B)/cp0."""
    return f"(({T}*{DBDT_CO2(T)})-({BCO2(T)}))/({CP0CO2(T)})"

def D2BDT2_CO2(T):
    """Analytic d2B/dT2 for the CO2 2nd virial fit [m3/(mol K2)]."""
    t = f"({T}/100)"
    return (f"0.0000000001*(-1131.240934/{t}^3+3608.789328/{t}^4"
            f"-24219.432384/{t}^5)")

def CPCO2(T, p):
    """Real-gas molar cp of CO2 [J/(mol K)]: cp0 - T*B''*p."""
    return f"({CP0CO2(T)}-{T}*{D2BDT2_CO2(T)}*{p})"

def CVCO2(T, p):
    """Real-gas molar cv of CO2 [J/(mol K)] from the 2nd virial."""
    rho = f"({p}/(({ZCO2(p,T)})*{RU}*{T}))"
    return (f"(({CP0CO2(T)})-{RU}-{T}*{RU}*(2*{DBDT_CO2(T)}"
            f"+{T}*{D2BDT2_CO2(T)})*{rho})")

def KH_CO2(T):
    """Henry's constant of CO2 in water [Pa]; IAPWS G7-04."""
    tr = f"({T}/647.096)"
    return (f"({PSAT(T)})*EXP(-8.55445/{tr}+4.01195*(1-{tr})^0.355/{tr}"
            f"+9.52345*{tr}^-0.41*EXP(1-{tr}))")

def B12(T):
    """Cross 2nd virial coefficient B12(H2O-CO2) [m3/mol]; fit to the
    Meyer & Harvey / Wheatley & Harvey data set."""
    return f"0.000001*(-236.26407+204871.53/{T}-59999434/{T}^2)"

def TDEW_HARDY(pv):
    """Hardy (1998) explicit rational inverse of p_sat [K]."""
    L = f"LN({pv})"
    return (f"(207.98233-20.156028*{L}+0.46778925*{L}^2-0.0000092288067*{L}^3)"
            f"/(1-0.13319669*{L}+0.0056577518*{L}^2-0.000075172865*{L}^3)")

def NEWTON_TDEW(Tn, psat_n, pv):
    """One Newton step on p_sat(T) = pv using the Clausius-Clapeyron slope."""
    return f"{Tn}-({psat_n}-{pv})*{RV}*{Tn}^2/(({psat_n})*{LV(Tn)})"

def KH_FROM_PSAT(psat_cell, T):
    """Henry's constant of CO2 in water [Pa] (IAPWS G7-04) from a p_sat cell."""
    tr = f"({T}/647.096)"
    return (f"({psat_cell})*EXP(-8.55445/{tr}+4.01195*(1-{tr})^0.355/{tr}"
            f"+9.52345*{tr}^(-0.41)*EXP(1-{tr}))")

def FONE(psat_cell, kh_cell, p, T, b12cell, use_f_cell=None):
    """Single-pass enhancement factor (seed x_v = p_sat/p), taking p_sat and
    k_H as cell references so the formula stays short. Accurate to 3.4e-5 in f
    versus the fully converged fixed point -- ample for sensitivity tables."""
    RT = f"({RU}*{T})"
    p, psat_cell, kh_cell = f"({p})", f"({psat_cell})", f"({kh_cell})"
    x0 = f"({psat_cell}/{p})"
    br = f"(({BWW(T)})*(2*{x0}-{x0}^2)+(2*({b12cell})-({BCO2(T)}))*(1-{x0})^2)"
    lnf = (f"(({BWW(T)})*{psat_cell}/{RT}+({VWL(T)})*({p}-{psat_cell})/{RT}"
           f"-{br}*{p}/{RT}+LN(1-({p}-{psat_cell})/{kh_cell}))")
    e = f"EXP({lnf})"
    return f"IF({use_f_cell}=1,{e},1)" if use_f_cell else e
